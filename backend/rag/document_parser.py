"""
文档解析：按扩展名识别 PDF / Word / 图片(OCR) / 纯文本，抽取为文本并分块

- .txt .md .csv   直接读取
- .pdf            pypdf 抽取文本
- .docx           python-docx 抽取段落 + 表格
- .png .jpg .bmp  rapidocr-onnxruntime 中文 OCR
"""
import os
import re
import logging
from typing import List, Tuple

logger = logging.getLogger(__name__)

TEXT_EXTS = {".txt", ".md", ".markdown", ".csv", ".json", ".log"}
PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".webp"}

_ocr_engine = None


def _get_ocr():
    """懒加载 RapidOCR（模型内置在 wheel 内，无网络下载）"""
    global _ocr_engine
    if _ocr_engine is None:
        from rapidocr_onnxruntime import RapidOCR
        _ocr_engine = RapidOCR()
    return _ocr_engine


def _read_text(filepath: str) -> str:
    with open(filepath, "r", encoding="utf-8") as f:
        return f.read()


def _read_pdf(filepath: str) -> str:
    from pypdf import PdfReader
    reader = PdfReader(filepath)
    pages = []
    for page in reader.pages:
        txt = page.extract_text()
        if txt:
            pages.append(txt)
    return "\n\n".join(pages)


def _read_docx(filepath: str) -> str:
    import docx
    doc = docx.Document(filepath)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # 表格单元格文本一并抽取
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_image_ocr(filepath: str) -> str:
    engine = _get_ocr()
    result, _ = engine(filepath)
    if not result:
        return ""
    return "\n".join(item[1] for item in result)


def parse_file(filepath: str) -> str:
    """按扩展名解析单个文件，返回纯文本"""
    ext = os.path.splitext(filepath)[1].lower()
    if ext in TEXT_EXTS:
        return _read_text(filepath)
    if ext in PDF_EXTS:
        return _read_pdf(filepath)
    if ext in DOCX_EXTS:
        return _read_docx(filepath)
    if ext in IMAGE_EXTS:
        return _read_image_ocr(filepath)
    raise ValueError(f"不支持的文件格式: {ext}")


def chunk_text(text: str, max_chars: int = 800, overlap: int = 100) -> List[str]:
    """按段落分块，超长段落硬切；保留 overlap 避免表格/语义被截断"""
    text = text.strip()
    if not text:
        return []
    paragraphs = [p.strip() for p in re.split(r"\n{2,}", text) if p.strip()]
    chunks: List[str] = []
    buf = ""
    for p in paragraphs:
        while len(p) > max_chars:
            if buf:
                chunks.append(buf)
                buf = ""
            chunks.append(p[:max_chars])
            p = p[max_chars - overlap:] if overlap else p[max_chars:]
        if not buf:
            buf = p
        elif len(buf) + len(p) + 2 <= max_chars:
            buf = f"{buf}\n\n{p}"
        else:
            chunks.append(buf)
            buf = p
    if buf:
        chunks.append(buf)
    return chunks


def ingest_files(filepaths: List[str], max_chars: int = 800, overlap: int = 100) -> List[Tuple[str, str]]:
    """解析多个文件并分块，返回 [(id, text), ...]，id 形如 '<文件名>#<块序号>'"""
    documents: List[Tuple[str, str]] = []
    for path in filepaths:
        if not os.path.exists(path):
            logger.warning(f"文件不存在，跳过: {path}")
            continue
        try:
            text = parse_file(path)
        except Exception as e:
            logger.warning(f"解析失败 {path}: {e}")
            continue
        name = os.path.basename(path)
        for i, chunk in enumerate(chunk_text(text, max_chars, overlap)):
            documents.append((f"{name}#{i}", chunk))
    logger.info(f"文档解析完成: {len(filepaths)} 个文件 → {len(documents)} 个块")
    return documents


if __name__ == "__main__":
    # 自检：分块逻辑正确（无外部文件依赖）
    sample = (
        "第一段：咖啡品类 Latte 现金支付 销量 38.7。\n\n"
        + "第二段：CCUS 项目 storage hub 钢铁行业 捕集能力 2 Mt。\n\n"
        + "长段落：" + "电动汽车 BEV 保有量 " * 80
    )
    chunks = chunk_text(sample, max_chars=200, overlap=20)
    assert chunks, "分块不应为空"
    assert all(len(c) <= 200 for c in chunks), "每块不得超过 max_chars"
    print(f"[OK] 分块自检通过: {len(chunks)} 块, 最大块长 {max(len(c) for c in chunks)}")
